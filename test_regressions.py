"""Offline behavior tests: each fixture owns its files, random source and I/O."""
import io
import random
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import main
from terms_exporter import export_docx, export_txt
from terms_loader import Term, load_terms, parse_terms, sample_terms
from test_glossary import FakeResponse, FakeSession


HTML = ('<table><tr><th>ru</th><th>kz</th></tr>'
        '<tr><td>name</td><td>атау</td></tr>'
        '<tr><td>description</td><td>анықтама</td></tr></table>')


class InputAndParsingTests(unittest.TestCase):
    def test_unknown_language_is_rejected(self):
        for operation in (lambda: parse_terms(HTML, 'xx'),
                          lambda: export_txt([], 'xx', '.'),
                          lambda: export_docx([], 'xx', '.')):
            with self.subTest(operation=operation), self.assertRaises(ValueError):
                operation()

    def test_duplicates_removed_and_case_insensitive_language(self):
        row = '<tr><td>name</td><td>атау</td></tr><tr><td>description</td><td>анықтама</td></tr>'
        self.assertEqual(parse_terms(HTML.replace('</table>', row + '</table>'), 'РУС'),
                         [Term('name', 'description')])

    def test_missing_translation_rejected(self):
        with self.assertRaises(ValueError):
            parse_terms('<table><tr></tr><tr><td>a</td></tr><tr><td>b</td></tr></table>', 'каз')

    def test_sampling_boundaries_reproducibility_and_input_unchanged(self):
        terms = [Term(str(n), 'value') for n in range(10)]
        snapshot = list(terms)
        self.assertEqual(sample_terms(terms, 5, random.Random(24)),
                         sample_terms(terms, 5, random.Random(24)))
        self.assertEqual(terms, snapshot)
        for count in (0, -1, 11):
            with self.subTest(count=count), self.assertRaises(ValueError):
                sample_terms(terms, count)
        with self.assertRaises(ValueError):
            sample_terms([], 1)


class CacheAndExportTests(unittest.TestCase):
    def test_valid_cache_avoids_network(self):
        with tempfile.TemporaryDirectory() as directory:
            cache = Path(directory) / 'cache.html'
            cache.write_text(HTML, encoding='utf-8')
            with patch('requests.get', side_effect=AssertionError('unexpected network')):
                self.assertEqual(load_terms('рус', cache), [Term('name', 'description')])

    def test_invalid_encoding_cache_is_replaced_only_after_validation(self):
        with tempfile.TemporaryDirectory() as directory:
            cache = Path(directory) / 'cache.html'
            cache.write_bytes(b'\xff')
            with patch('requests.get', return_value=FakeResponse([HTML.encode('utf-8')])):
                self.assertEqual(load_terms('каз', cache), [Term('атау', 'анықтама')])
            self.assertEqual(cache.read_text(encoding='utf-8'), HTML)

    def test_insecure_redirect_hop_rejected(self):
        response = FakeResponse([HTML.encode('utf-8')])
        hop = FakeResponse([])
        hop.url = 'http://example.invalid/redirect'
        response.history = [hop]
        with tempfile.TemporaryDirectory() as directory, self.assertRaises(ValueError):
            load_terms('рус', Path(directory) / 'cache.html', FakeSession(response))

    def test_failed_cache_replace_preserves_original_and_cleans_temporary(self):
        with tempfile.TemporaryDirectory() as directory:
            cache = Path(directory) / 'cache.html'
            cache.write_text('invalid', encoding='utf-8')
            with patch('terms_loader.os.replace', side_effect=OSError('locked')):
                with self.assertRaises(OSError):
                    load_terms('рус', cache, FakeSession(FakeResponse([HTML.encode('utf-8')])))
            self.assertEqual(cache.read_text(), 'invalid')
            self.assertEqual(list(Path(directory).iterdir()), [cache])

    def test_export_replace_failure_preserves_existing_file(self):
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / 'glossary_ru.txt'
            destination.write_text('original', encoding='utf-8')
            with patch('terms_exporter.os.replace', side_effect=OSError('locked')):
                with self.assertRaises(OSError):
                    export_txt([Term('a', 'b')], 'рус', directory)
            self.assertEqual(destination.read_text(), 'original')
            self.assertEqual(list(Path(directory).iterdir()), [destination])

    def test_kazakh_document_header_and_empty_export(self):
        import docx
        with tempfile.TemporaryDirectory() as directory:
            path = export_docx([Term('атау', 'анықтама')], 'каз', directory)
            table = docx.Document(path).tables[0]
            self.assertEqual(table.cell(0, 1).text, 'Терминнің мағынасы')
            self.assertEqual(table.cell(1, 1).text, 'анықтама')
            self.assertEqual(export_txt([], 'каз', directory).read_text(), '')


class CommandLineTests(unittest.TestCase):
    def test_invalid_inputs_retry_then_export_both_formats(self):
        terms = [Term('a', 'b')]
        with patch('builtins.input', side_effect=['oops', '0', '501', '1', 'invalid', '']), \
                patch('main.load_terms', return_value=terms), \
                patch('main.export_txt', return_value='out.txt') as txt, \
                patch('main.export_docx', return_value='out.docx') as docx, \
                patch('sys.stdout', new_callable=io.StringIO):
            self.assertEqual(main.main(), 0)
            self.assertEqual(txt.call_args.args[:2], (terms, 'рус'))
            self.assertEqual(docx.call_args.args[:2], (terms, 'рус'))

    def test_network_failure_returns_error_without_export(self):
        with patch('builtins.input', side_effect=['1', 'каз']), \
                patch('main.load_terms', side_effect=OSError('offline')), \
                patch('main.export_txt') as exporter, \
                patch('sys.stderr', new_callable=io.StringIO) as stderr:
            self.assertEqual(main.main(), 1)
            self.assertIn('offline', stderr.getvalue())
            exporter.assert_not_called()
